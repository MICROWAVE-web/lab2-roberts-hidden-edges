#!/usr/bin/env python3
"""Отчёт ЛР2 в формате титульного листа АлтГТУ."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "report_assets"
OUT = ROOT / "reports" / "Кованов_ПИ32_Компьютерная_графика_Лаб2.docx"
GITHUB = "https://github.com/MICROWAVE-web/lab2-roberts-hidden-edges"


def font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def p(
    doc,
    text="",
    *,
    size=14,
    bold=False,
    italic=False,
    center=False,
    indent=False,
    after=6,
    before=0,
):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    font(run, size=size, bold=bold, italic=italic)
    return para


def formula(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(text)
    font(run, name="Cambria Math", size=13, italic=True)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def code_block(doc, lines):
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.first_line_indent = Cm(0)
        run = para.add_run(line)
        font(run, name="Courier New", size=10)


def add_figure(doc, path: Path, caption: str, width_in=5.6):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    p(doc, caption, center=True, size=12, italic=True, after=10)


def bullet(doc, item: str):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    para.clear()
    run = para.add_run(item)
    font(run, size=14)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1.5)

    # --- Титул ---
    for line in [
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное бюджетное образовательное учреждение",
        "высшего образования",
        "«Алтайский государственный технический университет им. И. И. Ползунова»",
        "Факультет информационных технологий",
        "Кафедра прикладной математики",
    ]:
        p(doc, line, center=True, size=12, after=0)

    p(doc, "", after=20)
    p(doc, "Отчет защищен с оценкой _____", center=True, size=12, after=0)
    p(doc, "Преподаватель _____________", center=True, size=12, after=0)
    p(doc, "(подпись)", center=True, size=10, after=0)
    p(doc, "«___» ____________ 2026 г.", center=True, size=12, after=24)

    p(doc, "Отчет", center=True, bold=True, size=16, after=8)
    p(doc, "По лабораторной работе №2", center=True, bold=True, size=14, after=8)
    p(
        doc,
        "«Реализация алгоритмов удаления невидимых линий и поверхностей",
        center=True,
        bold=True,
        size=14,
        after=0,
    )
    p(
        doc,
        "при преобразовании изображения сложных пространственных сцен»",
        center=True,
        bold=True,
        size=14,
        after=8,
    )
    p(
        doc,
        "Вариант 2: алгоритм Робертса (невидимые рёбра выпуклого тела с динамикой)",
        center=True,
        size=12,
        after=8,
    )
    p(doc, "по дисциплине «Компьютерная графика»", center=True, size=14, after=20)
    p(doc, "Студент группы ПИ-32 Кованов Алексей Вадимович", center=True, size=14, after=4)
    p(doc, "Преподаватель Потапов Даниил Петрович", center=True, size=14, after=24)
    p(doc, "Барнаул 2026", center=True, size=14, after=12)

    doc.add_page_break()

    # --- Цель ---
    p(doc, "Цель работы", bold=True, size=14, after=6)
    p(
        doc,
        "Изучить алгоритмы удаления невидимых линий и поверхностей; программно реализовать "
        "простой алгоритм удаления невидимых рёбер выпуклого тела (алгоритм Робертса) "
        "с динамикой (вращение). Управление — мышь и клавиатура. "
        "Для демонстрации работы алгоритма выводятся текущие параметры классификации "
        "граней и рёбер.",
        indent=True,
    )

    p(doc, "Ссылка на исходный код программы:", bold=True, after=4)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(GITHUB)
    font(run, size=12, bold=True)
    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)

    p(
        doc,
        "Репозиторий GitHub: MICROWAVE-web / lab2-roberts-hidden-edges. "
        "Язык Python 3. Матрицы и проекции — вручную (модуль math3d.py из базы ЛР1). "
        "Готовые 3D-библиотеки не используются.",
        indent=True,
        after=10,
    )

    # --- 1 ---
    p(doc, "1. Постановка задачи", bold=True, size=14, before=8)
    p(
        doc,
        "При построении проекции трёхмерного объекта на плоскость экрана часть рёбер "
        "и граней оказывается невидимой (закрыта самим объектом). Необходимо "
        "сгенерировать 5–6 многоугольников (от 3 до 6 сторон) в пространстве и удалить "
        "невидимые части одним из методов.",
        indent=True,
    )
    p(
        doc,
        "Выбран вариант 2: простой алгоритм удаления невидимых рёбер у выпуклого тела "
        "с динамикой (вращение) — алгоритм Робертса. Допускается выполнение на базе "
        "преобразований из лабораторной работы №1.",
        indent=True,
    )

    # --- 2 ---
    p(doc, "2. Краткие теоретические сведения", bold=True, size=14, before=10)

    p(doc, "2.1. Выпуклое тело и ограничение метода", bold=True, size=14)
    p(
        doc,
        "Алгоритм Робертса в рассматриваемой форме использует свойство выпуклости: "
        "ребро невидимо тогда и только тогда, когда обе смежные с ним грани являются "
        "тыльными относительно наблюдателя. Объект лабораторной работы №1 (буква «А») "
        "невыпуклый, поэтому для варианта «Робертс» выбран выпуклый многогранник: "
        "прямоугольный параллелепипед (6 четырёхугольных граней) либо квадратная "
        "пирамида (5 граней: основание и 4 треугольных ската).",
        indent=True,
    )

    p(doc, "2.2. Шаги алгоритма Робертса", bold=True, size=14)
    p(doc, "1) Для каждой грани вычислить внешнюю нормаль:", indent=True)
    formula(doc, "n = normalize( (P₁ − P₀) × (P₂ − P₀) )")
    p(doc, "2) Классифицировать грань относительно наблюдателя eye:", indent=True)
    formula(doc, "to_eye = normalize(eye − c),   где c — центр грани")
    formula(doc, "грань лицевая (front), если n · to_eye > 0, иначе тыльная (back)")
    p(doc, "3) Классифицировать каждое ребро:", indent=True)
    bullet(doc, "ребро невидимо, если обе смежные грани тыльные;")
    bullet(
        doc,
        "ребро видимо, если хотя бы одна смежная грань лицевая (в том числе рёбра контура: front + back).",
    )

    p(doc, "2.3. Связь с аффинными преобразованиями (ЛР1)", bold=True, size=14)
    p(
        doc,
        "Вершины сначала преобразуются матрицами масштаба, поворота и переноса "
        "(однородные координаты, p′ = M·p), затем на преобразованных вершинах "
        "выполняется классификация Робертса, после чего видимые рёбра проецируются "
        "на экран (перспективная или ортогональная проекция).",
        indent=True,
    )

    # --- 3 ---
    p(doc, "3. Структура программы", bold=True, size=14, before=10)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate(
        [
            ("Файл", "Назначение"),
            ("roberts.py", "Нормали граней, классификация граней и рёбер (Робертс)"),
            ("polyhedron.py", "Модели выпуклых тел: параллелепипед и пирамида"),
            ("math3d.py", "Матрицы 4×4, перенос, поворот, масштаб, проекция (база ЛР1)"),
            ("main.py", "Окно, ввод, цикл отрисовки, панель демонстрации, анимация"),
        ]
    ):
        for j, val in enumerate((a, b)):
            cell = table.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=12, bold=(i == 0))

    p(doc, "Запуск:", indent=True, before=8)
    code_block(doc, ["/usr/local/bin/python3 main.py", "# или", "./run.sh"])

    # --- 4 ---
    p(doc, "4. Фрагменты кода", bold=True, size=14, before=12)
    p(doc, "Классификация грани и рёбер (roberts.py):", indent=True)
    code_block(
        doc,
        [
            "def is_front_face(vertices, face, eye):",
            "    n = face_normal(vertices, face)",
            "    c = face_center(vertices, face)",
            "    to_eye = normalize(sub(eye, c))",
            "    nd = dot(n, to_eye)",
            "    return nd > 1e-9, n, nd",
            "",
            "# Ребро видимо, если хотя бы одна смежная грань лицевая:",
            "visible = any(face_front[fi] for fi in adjacent_faces)",
        ],
    )
    p(doc, "Применение в кадре отрисовки (main.py):", indent=True, before=8)
    code_block(
        doc,
        [
            "pts = transform_points(world, self.vertices)",
            "face_front, edge_visible, face_info = roberts_classify(",
            "    pts, self.faces, self.edges, self.eye_world()",
            ")",
            "for (a, b), vis in zip(self.edges, edge_visible):",
            "    if vis: draw_solid(...)",
            "    elif self.show_hidden: draw_dashed(...)",
        ],
    )

    # --- 5 ---
    p(doc, "5. Управление программой", bold=True, size=14, before=12)
    ctrl = [
        ("Действие", "Управление"),
        ("Анимация вращения", "Пробел"),
        ("Скрытые рёбра (пунктир)", "H"),
        ("Нормали граней", "N"),
        ("Параллелепипед / пирамида", "B / P"),
        ("Проекция", "1 — ортогональная, 2 — перспективная"),
        ("Вращение мышью", "ЛКМ + движение"),
        ("Перенос / масштаб", "WASD, Q/E, колесо"),
        ("Сброс / выход", "0 / Esc"),
    ]
    t2 = doc.add_table(rows=len(ctrl), cols=2)
    t2.style = "Table Grid"
    for i, (a, b) in enumerate(ctrl):
        for j, val in enumerate((a, b)):
            cell = t2.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=12, bold=(i == 0))

    # --- 6 ---
    p(doc, "6. Результаты работы программы (скриншоты)", bold=True, size=14, before=12)
    p(
        doc,
        "На рисунках: красные сплошные линии — видимые рёбра; серый пунктир — скрытые "
        "(при включённом режиме H); зелёные/серые отрезки от центров граней — нормали.",
        indent=True,
    )
    figures = [
        ("fig1_roberts_hidden_on.png", "Рис. 1. Робертс: видимые рёбра и скрытые (пунктир), нормали"),
        ("fig2_roberts_hidden_off.png", "Рис. 2. Только видимые рёбра (скрытые отключены)"),
        ("fig3_rotated.png", "Рис. 3. Динамика: объект после поворота, пересчёт видимости"),
    ]
    for fname, cap in figures:
        path = ASSETS / fname
        if path.exists():
            add_figure(doc, path, cap)

    # --- 7 ---
    p(doc, "7. Демонстрация алгоритма", bold=True, size=14, before=8)
    p(
        doc,
        "Для защиты предусмотрена боковая панель: для каждой грани выводятся признак "
        "FRONT/back и значение скалярного произведения n·to_eye; для каждого ребра — "
        "VIS/hid. При вращении классификация обновляется каждый кадр, что наглядно "
        "показывает работу алгоритма Робертса.",
        indent=True,
    )

    # --- 8 ---
    p(doc, "8. Проверка (тестирование)", bold=True, size=14, before=8)
    for item in [
        "запуск окна, отображение выпуклого многогранника;",
        "при вращении меняется набор видимых/скрытых рёбер;",
        "клавиша H включает/выключает отрисовку скрытых рёбер пунктиром;",
        "клавиша N показывает нормали лицевых и тыльных граней;",
        "переключение B/P меняет модель (6 или 5 граней);",
        "панель справа согласована с картинкой (FRONT/back, VIS/hid);",
        "перенос, масштаб и смена проекции работают как в ЛР1.",
    ]:
        bullet(doc, item)

    p(
        doc,
        "Все проверки выполнены успешно: алгоритм стабильно классифицирует грани и рёбра "
        "при динамике и интерактивном управлении.",
        indent=True,
        before=6,
    )

    # --- Вывод ---
    p(doc, "Вывод", bold=True, size=14, before=12)
    p(
        doc,
        "В ходе лабораторной работы изучен и программно реализован алгоритм Робертса "
        "удаления невидимых рёбер выпуклого многогранника. Реализованы классификация "
        "граней (лицевые/тыльные) и рёбер (видимые/невидимые), динамическое вращение, "
        "панель демонстрации параметров алгоритма и визуализация скрытых рёбер и нормалей. "
        "Аффинные преобразования и проекции выполнены на базе модуля math3d.py лабораторной "
        f"работы №1 без готовых 3D-библиотек. Исходный код: {GITHUB}.",
        indent=True,
    )

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
